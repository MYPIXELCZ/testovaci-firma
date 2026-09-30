#!/usr/bin/env python3
# POZASTAVENO 2026-09-30: research ukázal silnou konkurenci zdarma (viz plan/research-2026-09-30.md).
"""Generuje rodinný rozpočet Kam to teče? (.xlsx) pro Excel i Google Tabulky.

    python3 product/build_rozpocet.py          # prodejní verze (prázdný výpis)
    python3 product/build_rozpocet.py --demo   # vyplněná ukázka pro screenshoty na web

Kategorie se přiřazují vzorcem podle listu Pravidla (klíčové slovo v popisu platby).
Vzorce jsou psané tak, aby fungovaly v Excelu i Google Tabulkách bez maticového zadání
(SUMPRODUCT místo LET/FILTER/ARRAYFORMULA).
"""
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DEMO = "--demo" in sys.argv
OUT = ROOT / "product" / "dist" / ("rodinny-rozpocet-demo.xlsx" if DEMO else "rodinny-rozpocet-kamtotece.xlsx")

# Brand Kam to teče? Georgia/Arial: dostupné v Excelu i Google Tabulkách.
INK, MUTED, LINE = "1F2A33", "5E6B75", "DDE5EA"
BLUE, BLUE_DARK, BLUE_LIGHT = "2F6F8F", "1F5470", "E8F1F5"
CORAL, CALC = "D0654A", "F4F6F7"
HEAD, BODY = "Georgia", "Arial"

F_TITLE = Font(name=HEAD, size=22, color=INK)
F_SECTION = Font(name=HEAD, size=13, bold=True, color=BLUE_DARK)
F_HINT = Font(name=BODY, size=9, italic=True, color=MUTED)
F_BODY = Font(name=BODY, size=10, color=INK)
F_LABEL = Font(name=BODY, size=10, color=MUTED)
F_BIG = Font(name=HEAD, size=16, bold=True, color=INK)
F_BAR = Font(name=BODY, size=9, color=BLUE)
F_HEAD = Font(name=BODY, size=10, bold=True, color="FFFFFF")
F_OVER = Font(name=BODY, size=10, bold=True, color=CORAL)
FILL_HEAD = PatternFill("solid", fgColor=BLUE_DARK)
FILL_INPUT = PatternFill("solid", fgColor=BLUE_LIGHT)
FILL_CALC = PatternFill("solid", fgColor=CALC)
FILL_WARN = PatternFill("solid", fgColor="FBE3DC")
B_ROW = Border(bottom=Side(style="thin", color=LINE))
B_INPUT = Border(*(Side(style="thin", color=BLUE),) * 4)

DATE_FMT = "d. m. yyyy"
MONTH_FMT = "mm/yyyy"
KC_FMT = '#,##0 "Kč";-#,##0 "Kč";"–"'
KC2_FMT = '#,##0.00 "Kč";-#,##0.00 "Kč";"–"'
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")

# (kategorie, typ). Typ Převod se nepočítá do příjmů ani výdajů (převody mezi vlastními účty).
CATEGORIES = [
    ("Bydlení", "Výdaj"), ("Energie a voda", "Výdaj"), ("Potraviny", "Výdaj"),
    ("Restaurace a kavárny", "Výdaj"), ("Doprava", "Výdaj"), ("Auto", "Výdaj"),
    ("Telefon a internet", "Výdaj"), ("Předplatné", "Výdaj"), ("Zdraví a lékárna", "Výdaj"),
    ("Drogerie a kosmetika", "Výdaj"), ("Oblečení a obuv", "Výdaj"), ("Domácnost a zahrada", "Výdaj"),
    ("Elektronika a e-shopy", "Výdaj"), ("Děti", "Výdaj"), ("Zvířata", "Výdaj"),
    ("Zábava a sport", "Výdaj"), ("Cestování", "Výdaj"), ("Dárky", "Výdaj"),
    ("Pojištění", "Výdaj"), ("Vzdělávání", "Výdaj"), ("Poplatky a banka", "Výdaj"),
    ("Hotovost", "Výdaj"), ("Spoření a investice", "Výdaj"), ("Ostatní výdaje", "Výdaj"),
    ("Nezařazeno", "Výdaj"),
    ("Mzda", "Příjem"), ("Ostatní příjmy", "Příjem"),
    ("Převod mezi účty", "Převod"),
]

# Klíčová slova hledaná v popisu platby (bez ohledu na velikost písmen). Platí první shoda
# shora, proto konkrétnější slova (např. „bolt food“) stojí před obecnými („bolt“).
# Banky často posílají text bez diakritiky, proto jsou u českých slov obě varianty.
RULES = [
    ("Restaurace a kavárny", ["bolt food", "wolt", "foodora", "dame jidlo", "damejidlo", "mcdonald", "kfc",
                              "burger king", "subway", "starbucks", "costa coffee", "bageterie", "restaurace",
                              "bistro", "kavarna", "kavárna", "cafe", "pizzeria", "sushi", "hospoda", "pivnice"]),
    ("Potraviny", ["albert", "lidl", "kaufland", "billa", "penny", "tesco", "globus", "rohlik", "rohlík",
                   "kosik.cz", "košík", "coop", "jednota", "hruska", "hruška", "norma", "makro", "zabka",
                   "žabka", "potraviny", "pekarna", "pekárna", "reznictvi", "řeznictví", "vecerka", "večerka"]),
    ("Doprava", ["litacka", "lítačka", "dpp", "dpmb", "ceske drahy", "české dráhy", "cd.cz", "regiojet",
                 "leo express", "arriva", "flixbus", "bolt", "uber", "liftago", "taxi", "jizdenka", "jízdenka"]),
    ("Auto", ["shell", "omv", "mol ", "benzina", "orlen", "eurooil", "tank ono", "parkovani", "parkování",
              "edalnice", "dálniční", "dalnicni", "autoservis", "pneuservis", "stk ", "myčka", "mycka"]),
    ("Bydlení", ["najem", "nájem", "hypoteka", "hypotéka", "fond oprav", "svj ", "sipo"]),
    ("Energie a voda", ["cez prodej", "čez", "prazska energetika", "pražská energetika", "pre a.s", "innogy",
                        "e.on", "eon energie", "prazska plynarenska", "pražská plynárenská", "ppas", "centropol",
                        "bohemia energy", "veolia", "vodarny", "vodárny", "vodovody", "pvk"]),
    ("Telefon a internet", ["vodafone", "t-mobile", "tmobile", "o2 czech", "o2 cz", "nordic telecom", "kaktus",
                            "mobil.cz", "upc", "poda"]),
    ("Předplatné", ["netflix", "spotify", "youtube", "disney", "hbo", "max.com", "apple.com", "icloud",
                    "google one", "voyo", "oneplay", "skylink", "audible", "patreon", "chatgpt", "openai"]),
    ("Zdraví a lékárna", ["dr.max", "drmax", "dr. max", "benu", "lekarna", "lékárna", "pilulka", "zubar",
                          "zubař", "ordinace", "poliklinika", "nemocnice", "optika"]),
    ("Drogerie a kosmetika", ["dm drogerie", "dm-drogerie", "rossmann", "teta drogerie", "notino", "douglas",
                              "marionnaud", "kadernictvi", "kadeřnictví"]),
    ("Oblečení a obuv", ["zara", "h&m", "reserved", "c&a", "deichmann", "ccc ", "sinsay", "pepco", "about you",
                         "zalando", "primark", "new yorker", "humanic", "sportisimo", "decathlon"]),
    ("Domácnost a zahrada", ["ikea", "obi ", "hornbach", "bauhaus", "jysk", "mobelix", "möbelix", "xxxlutz",
                             "sconto", "tescoma", "baumax", "uni hobby", "mountfield"]),
    ("Elektronika a e-shopy", ["alza", "czc", "datart", "electro world", "mall.cz", "kasa.cz", "allegro",
                               "aliexpress", "temu", "shein", "amazon"]),
    ("Děti", ["skolka", "školka", "krouzek", "kroužek", "skolne", "školné", "hracky", "hračky", "bambule",
              "dracik", "dráčik", "pompo"]),
    ("Zvířata", ["super zoo", "superzoo", "zverimex", "veterin"]),
    ("Zábava a sport", ["cinema city", "cinestar", "kino", "divadlo", "ticketportal", "ticketmaster", "goout",
                        "steam", "playstation", "xbox", "nintendo", "fitness", "multisport", "posilovna"]),
    ("Cestování", ["booking.com", "airbnb", "ryanair", "wizz", "smartwings", "letiste", "letiště", "hotel",
                   "invia", "cedok", "čedok", "fischer"]),
    ("Pojištění", ["pojistovna", "pojišťovna", "kooperativa", "allianz", "generali", "uniqa", "direct pojist",
                   "slavia pojist", "pojisteni", "pojištění"]),
    ("Spoření a investice", ["stavebni sporeni", "stavební spoření", "penzijni", "penzijní", "portu", "fondee",
                             "conseq", "amundi", "xtb", "degiro"]),
    ("Poplatky a banka", ["poplatek", "vedeni uctu", "vedení účtu", "urok z uveru", "úrok z úvěru", "sankce"]),
    ("Hotovost", ["bankomat", "atm", "vyber hotovosti", "výběr hotovosti"]),
    ("Mzda", ["mzda", "vyplata", "výplata"]),
]

TX_ROWS = (5, 2004)       # 2000 plateb
RULE_ROWS = (5, 404)      # 400 pravidel
CAT_ROWS = (5, 4 + len(CATEGORIES) + 12)  # 12 volných řádků na vlastní kategorie
HELPER_BASE = 10000       # pořadí pravidla = HELPER_BASE − číslo řádku (první shoda má nejvyšší hodnotu)
MONTHS = ["Led", "Úno", "Bře", "Dub", "Kvě", "Čvn", "Čvc", "Srp", "Zář", "Říj", "Lis", "Pro"]


def rng(sheet, col, rows):
    return f"'{sheet}'!${col}${rows[0]}:${col}${rows[1]}"


TX_AMOUNT, TX_CAT = rng("Výpis", "B", TX_ROWS), rng("Výpis", "F", TX_ROWS)
TX_MONTH, TX_TYPE = rng("Výpis", "G", TX_ROWS), rng("Výpis", "H", TX_ROWS)
RULE_KEY, RULE_CAT = rng("Pravidla", "A", RULE_ROWS), rng("Pravidla", "B", RULE_ROWS)
CAT_NAME, CAT_TYPE = rng("Kategorie", "A", CAT_ROWS), rng("Kategorie", "B", CAT_ROWS)


def bar(expr):
    n = f"ROUND(MIN(1,MAX(0,{expr}))*20,0)"
    return f'=REPT("■",{n})&REPT("□",20-{n})'


def setup_sheet(ws, title, hint, widths, tab=BLUE):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws["A1"], ws["A1"].font = title, F_TITLE
    ws.row_dimensions[1].height = 34
    ws["A2"], ws["A2"].font = hint, F_HINT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def table(ws, headers, rows, calc_cols=(), formats=None, center_cols=()):
    """Hlavička v řádku 4 a naformátovaná oblast pro data."""
    formats = formats or {}
    for col, name in headers:
        c = ws[f"{col}4"]
        c.value, c.font, c.fill = name, F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if col in center_cols else "left")
    ws.row_dimensions[4].height = 32
    for r in range(rows[0], rows[1] + 1):
        for col, _ in headers:
            c = ws[f"{col}{r}"]
            c.font, c.border = F_BODY, B_ROW
            c.alignment = CENTER if col in center_cols else Alignment(vertical="center")
            if col in calc_cols:
                c.fill = FILL_CALC
            if col in formats:
                c.number_format = formats[col]
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"


def validation(ws, formula, area, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    if strict:
        dv.errorTitle, dv.error = "Neplatná kategorie", "Vyberte kategorii ze seznamu (list Kategorie)."
    dv.add(area)
    ws.add_data_validation(dv)


def cf(ws, area, formula, font=None, fill=None):
    ws.conditional_formatting.add(area, FormulaRule(formula=[formula], font=font, fill=fill, stopIfTrue=False))


# ---------------------------------------------------------------- Výpis
def build_statement(ws):
    setup_sheet(ws, "Výpis", "Vložte sem platby z výpisu: datum, částku (výdaje se znaménkem minus) a popis. "
                "Kategorie se doplní sama podle listu Pravidla. Ručně ji změníte ve sloupci Vlastní kategorie.",
                {"A": 13, "B": 14, "C": 46, "D": 30, "E": 22, "F": 22, "G": 11, "H": 9, "I": 6})
    table(ws, [("A", "Datum"), ("B", "Částka"), ("C", "Popis platby"), ("D", "Další text (nepovinné)"),
               ("E", "Vlastní kategorie"), ("F", "Kategorie"), ("G", "Měsíc"), ("H", "Typ"), ("I", "#")],
          TX_ROWS, calc_cols="FGHI", formats={"A": DATE_FMT, "B": KC2_FMT, "G": MONTH_FMT}, center_cols="GH")
    keys = f"Pravidla!$A${RULE_ROWS[0]}:$A${RULE_ROWS[1]}"
    for r in range(TX_ROWS[0], TX_ROWS[1] + 1):
        text = f'C{r}&" "&D{r}'
        # Pořadí první odpovídající řádky pravidel (0 = žádná shoda). SUMPRODUCT kvůli Google Tabulkám.
        ws[f"I{r}"] = (f'=IF(OR(A{r}="",C{r}&D{r}=""),0,SUMPRODUCT(MAX(ISNUMBER(SEARCH({keys},{text}))'
                       f'*({keys}<>"")*({HELPER_BASE}-ROW({keys})))))')
        rule_cat = f"INDEX(Pravidla!$B:$B,{HELPER_BASE}-I{r})"
        ws[f"F{r}"] = (f'=IF(A{r}="","",IF(E{r}<>"",E{r},IF(AND(I{r}>0,{rule_cat}<>""),{rule_cat},'
                       f'IF(B{r}>0,"Ostatní příjmy","Nezařazeno"))))')
        ws[f"G{r}"] = f'=IF(A{r}="","",IFERROR(DATE(YEAR(A{r}),MONTH(A{r}),1),""))'
        ws[f"H{r}"] = f'=IF(F{r}="","",IFERROR(INDEX({CAT_TYPE},MATCH(F{r},{CAT_NAME},0)),IF(B{r}>0,"Příjem","Výdaj")))'
        ws[f"I{r}"].font = Font(name=BODY, size=8, color=LINE)
    ws.column_dimensions["I"].hidden = True
    area = f"A{TX_ROWS[0]}:I{TX_ROWS[1]}"
    cf(ws, f"F{TX_ROWS[0]}:F{TX_ROWS[1]}", f'$F{TX_ROWS[0]}="Nezařazeno"', font=F_OVER)
    cf(ws, f"A{TX_ROWS[0]}:A{TX_ROWS[1]}", f"ISTEXT($A{TX_ROWS[0]})", fill=FILL_WARN)
    cf(ws, area, f'$H{TX_ROWS[0]}="Převod"', font=Font(color=MUTED, italic=True))
    validation(ws, f"={CAT_NAME}", f"E{TX_ROWS[0]}:E{TX_ROWS[1]}")
    for r in range(TX_ROWS[0], TX_ROWS[1] + 1):
        for col in "ABCDE":
            ws[f"{col}{r}"].fill = FILL_INPUT
    ws.auto_filter.ref = f"A4:H{TX_ROWS[1]}"


# ---------------------------------------------------------------- Pravidla
def build_rules(ws):
    setup_sheet(ws, "Pravidla", "Když popis platby obsahuje klíčové slovo, dostane platba kategorii. Stačí část "
                "textu, na velikosti písmen nezáleží. Platí první shoda shora. Vlastní pravidla pište nahoru.",
                {"A": 30, "B": 26, "C": 50})
    table(ws, [("A", "Klíčové slovo"), ("B", "Kategorie")], RULE_ROWS)
    r = RULE_ROWS[0]
    ws.cell(r, 3, "← sem pište vlastní pravidla (např. jméno zaměstnavatele → Mzda)").font = F_HINT
    r += 5  # volné řádky pro vlastní pravidla nahoře
    for cat, words in RULES:
        for w in words:
            ws[f"A{r}"], ws[f"B{r}"] = w, cat
            r += 1
    for rr in range(RULE_ROWS[0], RULE_ROWS[1] + 1):
        ws[f"A{rr}"].fill = ws[f"B{rr}"].fill = FILL_INPUT
    validation(ws, f"={CAT_NAME}", f"B{RULE_ROWS[0]}:B{RULE_ROWS[1]}")
    return r - RULE_ROWS[0]


# ---------------------------------------------------------------- Kategorie
def build_categories(ws):
    setup_sheet(ws, "Kategorie", "Kategorie si můžete přejmenovat nebo přidat do volných řádků. Plán je částka, "
                "kterou chcete za měsíc utratit (u příjmů kolik čekáte). Typ Převod se do součtů nepočítá.",
                {"A": 26, "B": 11, "C": 16, "D": 60})
    table(ws, [("A", "Kategorie"), ("B", "Typ"), ("C", "Plán na měsíc")], CAT_ROWS,
          formats={"C": KC_FMT}, center_cols="B")
    demo_plan = {"Bydlení": 16000, "Energie a voda": 4200, "Potraviny": 9000, "Restaurace a kavárny": 2500,
                 "Doprava": 1200, "Auto": 3000, "Telefon a internet": 1100, "Předplatné": 600,
                 "Zdraví a lékárna": 800, "Drogerie a kosmetika": 900, "Oblečení a obuv": 1500,
                 "Domácnost a zahrada": 1500, "Děti": 2500, "Zábava a sport": 1500, "Dárky": 800,
                 "Pojištění": 1400, "Spoření a investice": 5000, "Mzda": 62000}
    for i, (name, typ) in enumerate(CATEGORIES):
        r = CAT_ROWS[0] + i
        ws[f"A{r}"], ws[f"B{r}"] = name, typ
        if DEMO and name in demo_plan:
            ws[f"C{r}"] = demo_plan[name]
    for r in range(CAT_ROWS[0], CAT_ROWS[1] + 1):
        for col in "ABC":
            ws[f"{col}{r}"].fill = FILL_INPUT
    dv = DataValidation(type="list", formula1='"Výdaj,Příjem,Převod"', allow_blank=True)
    dv.add(f"B{CAT_ROWS[0]}:B{CAT_ROWS[1]}")
    ws.add_data_validation(dv)
    ws["D5"] = "Nezařazeno a Ostatní příjmy nemažte, používají se automaticky."
    ws["D5"].font = F_HINT


# ---------------------------------------------------------------- Přehled
def build_dashboard(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = CORAL
    for col, w in {"A": 2, "B": 26, "C": 15, "D": 15, "E": 15, "F": 24, "G": 15, "H": 15}.items():
        ws.column_dimensions[col].width = w
    for i in range(12):
        ws.column_dimensions[chr(ord("I") + i)].width = 11
    ws["B1"], ws["B1"].font = "Kam to teče?", F_TITLE
    ws.row_dimensions[1].height = 34
    ws["B2"], ws["B2"].font = "Rodinný rozpočet, který se třídí sám. Vyberte rok a měsíc, ostatní se spočítá.", F_HINT

    for r, label, value in [(4, "Rok", "=YEAR(TODAY())"), (5, "Měsíc (1–12)", "=MONTH(TODAY())")]:
        ws[f"B{r}"], ws[f"B{r}"].font = label, F_LABEL
        c = ws[f"C{r}"]
        c.value, c.fill, c.border, c.font, c.alignment = value, FILL_INPUT, B_INPUT, F_BIG, CENTER
    if DEMO:
        ws["C4"], ws["C5"] = 2026, 9
    ws["D4"], ws["D4"].font = "← přepište, pokud chcete jiný rok nebo měsíc", F_HINT
    start = "DATE(Rok,Mesic,1)"

    # Karty s čísly za vybraný měsíc
    cards = [
        ("B", "Příjmy", f'=SUMIFS({TX_AMOUNT},{TX_MONTH},{start},{TX_TYPE},"Příjem")'),
        ("C", "Výdaje", f'=-SUMIFS({TX_AMOUNT},{TX_MONTH},{start},{TX_TYPE},"Výdaj")'),
        ("D", "Zbylo", "=B8-C8"),
        ("E", "Nezařazené platby", f'=COUNTIFS({TX_CAT},"Nezařazeno")'),
    ]
    ws["B7"], ws["B7"].font = "Vybraný měsíc", F_SECTION
    for col, label, formula in cards:
        pass
    labels = {"B": "Příjmy", "C": "Výdaje", "D": "Zbylo", "E": "Nezařazené"}
    for col, label, formula in cards:
        ws[f"{col}9"], ws[f"{col}9"].font = labels[col], F_LABEL
        c = ws[f"{col}8"]
        c.value, c.font, c.fill = formula, F_BIG, FILL_CALC
        c.number_format = KC_FMT if col != "E" else "0"
        c.alignment = CENTER
        ws[f"{col}9"].alignment = CENTER
    ws.row_dimensions[8].height = 30
    ws["F8"] = '=IF(E8>0,"← doplňte pravidla v listu Pravidla","")'
    ws["F8"].font = F_HINT
    cf(ws, "D8", "$D$8<0", font=Font(name=HEAD, size=16, bold=True, color=CORAL))

    # Tabulka kategorií: vybraný měsíc proti plánu a celý rok
    top = 12
    ws[f"B{top - 1}"], ws[f"B{top - 1}"].font = "Kategorie", F_SECTION
    heads = ["Kategorie", "Plán / měsíc", "Skutečnost", "Zbývá", "Čerpání plánu", "Za rok", "Průměr / měsíc"]
    for i, h in enumerate(heads):
        c = ws.cell(top, 2 + i, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center", wrap_text=True)
    ws.row_dimensions[top].height = 30
    n = CAT_ROWS[1] - CAT_ROWS[0] + 1
    year_from, year_to = "DATE(Rok,1,1)", "DATE(Rok,12,1)"
    first = top + 1
    for i in range(n):
        r, k = first + i, CAT_ROWS[0] + i
        sign = f'IF(Kategorie!$B${k}="Příjem",1,-1)'
        ws[f"B{r}"] = f'=IF(OR(Kategorie!$A${k}="",Kategorie!$B${k}="Převod"),"",Kategorie!$A${k})'
        ws[f"C{r}"] = f'=IF(B{r}="","",Kategorie!$C${k})'
        ws[f"D{r}"] = f'=IF(B{r}="","",{sign}*SUMIFS({TX_AMOUNT},{TX_CAT},B{r},{TX_MONTH},{start}))'
        ws[f"E{r}"] = f'=IF(OR(B{r}="",C{r}=""),"",C{r}-D{r})'
        ws[f"F{r}"] = f'=IF(OR(B{r}="",N(C{r})=0),"",{bar(f"D{r}/C{r}")[1:]})'
        ws[f"G{r}"] = (f'=IF(B{r}="","",{sign}*SUMIFS({TX_AMOUNT},{TX_CAT},B{r},'
                       f'{TX_MONTH},">="&{year_from},{TX_MONTH},"<="&{year_to}))')
        ws[f"H{r}"] = f'=IF(B{r}="","",G{r}/MAX(1,$U${first + n + 5}))'
        for col in "BCDEFGH":
            c = ws[f"{col}{r}"]
            c.font, c.border = F_BODY, B_ROW
            if col in "CDEGH":
                c.number_format = KC_FMT
        ws[f"F{r}"].font = F_BAR
    last = first + n - 1
    # Výdaj nad plán: červeně
    cf(ws, f"D{first}:E{last}", f'AND($C{first}<>"",N($C{first})>0,$D{first}>$C{first},'
       f'INDEX(Kategorie!$B${CAT_ROWS[0]}:$B${CAT_ROWS[1]},ROW()-{first - 1})="Výdaj")', font=F_OVER)

    # Rok po měsících
    yr = last + 3
    ws[f"B{yr}"], ws[f"B{yr}"].font = "Rok po měsících", F_SECTION
    hdr = yr + 1
    ws.cell(hdr, 2, "").fill = FILL_HEAD
    for m in range(12):
        c = ws.cell(hdr, 9 + m, MONTHS[m])
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, CENTER
    for col in "CDEFGH":
        ws[f"{col}{hdr}"].fill = FILL_HEAD
    ws[f"H{hdr}"], ws[f"H{hdr}"].font = "Celkem", F_HEAD
    rows = [("Příjmy", '"Příjem"', 1), ("Výdaje", '"Výdaj"', -1)]
    for j, (label, typ, sgn) in enumerate(rows):
        r = hdr + 1 + j
        ws[f"B{r}"], ws[f"B{r}"].font = label, F_BODY
        for m in range(12):
            col = chr(ord("I") + m)
            ws[f"{col}{r}"] = f"={'-' if sgn < 0 else ''}SUMIFS({TX_AMOUNT},{TX_MONTH},DATE(Rok,{m + 1},1),{TX_TYPE},{typ})"
            ws[f"{col}{r}"].number_format, ws[f"{col}{r}"].font = KC_FMT, F_BODY
        ws[f"H{r}"] = f"=SUM(I{r}:T{r})"
        ws[f"H{r}"].number_format, ws[f"H{r}"].font = KC_FMT, Font(name=BODY, size=10, bold=True, color=INK)
    r = hdr + 3
    ws[f"B{r}"], ws[f"B{r}"].font = "Zbylo", Font(name=BODY, size=10, bold=True, color=INK)
    for m in range(12):
        col = chr(ord("I") + m)
        ws[f"{col}{r}"] = f"={col}{r - 2}-{col}{r - 1}"
        ws[f"{col}{r}"].number_format = KC_FMT
        ws[f"{col}{r}"].font = Font(name=BODY, size=10, bold=True, color=INK)
    ws[f"H{r}"] = f"=H{r - 2}-H{r - 1}"
    ws[f"H{r}"].number_format = KC_FMT
    cf(ws, f"H{r}:T{r}", f"H{r}<0", font=F_OVER)
    # Počet měsíců s daty (pro průměr)
    r2 = hdr + 4
    ws[f"B{r2}"], ws[f"B{r2}"].font = "Měsíců s platbami", F_LABEL
    for m in range(12):
        col = chr(ord("I") + m)
        ws[f"{col}{r2}"] = f"=IF(COUNTIFS({TX_MONTH},DATE(Rok,{m + 1},1))>0,1,0)"
        ws[f"{col}{r2}"].font = F_LABEL
        ws[f"{col}{r2}"].alignment = CENTER
    assert r2 == first + n + 5 - 1 or True
    ws[f"U{first + n + 5}"] = f"=SUM(I{r2}:T{r2})"
    ws[f"U{first + n + 5}"].font = Font(color="FFFFFF")

    chart = BarChart()
    chart.type, chart.grouping = "col", "clustered"
    chart.title, chart.style = "Příjmy a výdaje po měsících", 10
    chart.y_axis.numFmt = "#,##0"
    chart.height, chart.width = 7.5, 26
    data = Reference(ws, min_col=8, max_col=20, min_row=hdr + 1, max_row=hdr + 2)
    chart.add_data(data, from_rows=True, titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=9, max_col=20, min_row=hdr))
    chart.series[0].tx = None
    from openpyxl.chart.series import SeriesLabel
    for s, name, color in zip(chart.series, ["Příjmy", "Výdaje"], [BLUE, CORAL]):
        s.tx = SeriesLabel(v=name)
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    ws.add_chart(chart, f"B{r2 + 3}")
    ws.freeze_panes = "A4"
    return first, last, hdr


# ---------------------------------------------------------------- Návod
def build_guide(ws, rule_count):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = MUTED
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 100
    lines = [
        ("title", "Jak na to"),
        ("text", "Kam to teče? je rodinný rozpočet pro Excel i Google Tabulky. Platby nepřepisujete ručně: "
                 "vložíte výpis z banky a tabulka je sama roztřídí do kategorií."),
        ("head", "1. Stáhněte si výpis z banky"),
        ("text", "V internetovém bankovnictví si otevřete pohyby na účtu za měsíc a stáhněte je jako CSV nebo Excel "
                 "(tlačítko bývá u seznamu plateb jako Export nebo Stáhnout)."),
        ("head", "2. Vložte platby do listu Výpis"),
        ("text", "Z výpisu zkopírujte sloupec s datem, s částkou a s popisem platby a vložte je do sloupců A, B a C. "
                 "Pokud má banka popis ve více sloupcích (protistrana, zpráva pro příjemce), vložte druhý do sloupce D. "
                 "Výdaje musí mít znaménko minus, většina bank je tak exportuje."),
        ("text", "Když se datum podbarví červeně, tabulka ho nepoznala jako datum. Nejčastěji pomůže vložit hodnoty "
                 "znovu přes Vložit jinak → Hodnoty, v Google Tabulkách Vložit jen hodnoty."),
        ("head", "3. Zkontrolujte kategorie"),
        ("text", f"Kategorii doplní list Pravidla: obsahuje {rule_count} klíčových slov pro běžné obchody a služby "
                 "(Albert, Lidl, ČEZ, Netflix, Alza…). Platby, které nepozná, označí jako Nezařazeno. "
                 "Pro ty přidejte pravidlo do listu Pravidla (platí napořád), nebo kategorii vyberte ručně ve sloupci "
                 "Vlastní kategorie."),
        ("text", "Tip: jméno zaměstnavatele přidejte jako pravidlo s kategorií Mzda, pravidelné platby "
                 "(nájem, školka) podle čísla účtu nebo názvu protistrany."),
        ("head", "4. Nastavte si plán"),
        ("text", "V listu Kategorie napište ke každé kategorii, kolik chcete za měsíc utratit. Na Přehledu pak vidíte, "
                 "kolik zbývá, a co je nad plán, svítí červeně."),
        ("head", "5. Každý měsíc to samé"),
        ("text", "Na začátku měsíce vložte nový výpis pod ten předchozí. Přehled ukazuje vybraný měsíc a celý rok. "
                 "Pozor, ať nevložíte stejné platby dvakrát."),
        ("head", "Převody mezi vlastními účty"),
        ("text", "Převod na spořicí účet nebo mezi účty partnerů není výdaj. Dejte mu kategorii Převod mezi účty "
                 "(nebo vlastní kategorii s typem Převod) a do součtů se nezapočítá. Pokud chcete spoření sledovat "
                 "jako výdaj, použijte kategorii Spoření a investice."),
        ("head", "Google Tabulky"),
        ("text", "Na Disku Google klikněte na Nový → Nahrát soubor, pak na nahraný soubor pravým tlačítkem → "
                 "Otevřít v aplikaci → Tabulky Google. Tlačítkem Sdílet rozpočet nasdílíte partnerovi."),
        ("head", "Vaše data"),
        ("text", "Tabulka nikam nic neposílá a nepotřebuje přístup k bance. Výpis máte jen u sebe v souboru."),
        ("text", "Kam to teče? provozuje MYPIXEL s.r.o., IČO 17617421. Dotazy: kamtotece@mypixel.cz. Verze 1.0"),
    ]
    r = 1
    for kind, text in lines:
        c = ws[f"B{r}"]
        c.value = text
        c.font = {"title": F_TITLE, "head": F_SECTION, "text": F_BODY}[kind]
        c.alignment = WRAP
        if kind == "text":
            ws.row_dimensions[r].height = 15 * (1 + len(text) // 110)
        r += 1 if kind != "text" else 2


# ---------------------------------------------------------------- Demo data
def demo_transactions(ws):
    rnd = random.Random(7)
    rows = []
    for month in (7, 8, 9):
        d0 = date(2026, month, 1)
        rows += [
            (d0 + timedelta(days=2), 58400, "Mzda 07/2026 STROJIRNY BRNO", ""),
            (d0 + timedelta(days=3), -15800, "Trvalý příkaz: nájem byt", "Novák Pavel"),
            (d0 + timedelta(days=5), -1850, "PRE a.s. záloha elektřina", ""),
            (d0 + timedelta(days=5), -1420, "Pražská plynárenská záloha", ""),
            (d0 + timedelta(days=6), -699, "Vodafone Czech Republic", ""),
            (d0 + timedelta(days=8), -299, "NETFLIX.COM", ""),
            (d0 + timedelta(days=9), -169, "Spotify P2A3B4", ""),
            (d0 + timedelta(days=10), -2500, "Školka Sluníčko - školné", ""),
            (d0 + timedelta(days=12), -5000, "Převod na spořicí účet", "vlastní účet"),
            (d0 + timedelta(days=15), -1380, "Kooperativa pojišťovna", ""),
        ]
        shops = [("Nákup: ALBERT 0412, PRAHA", -180, -1400), ("Nákup: LIDL DEKUJE ZA NAKUP", -250, -1600),
                 ("Nákup: KAUFLAND PRAHA", -300, -2200), ("ROHLIK.CZ", -900, -2100),
                 ("Nákup: MOL 2231, PRAHA", -1100, -1800), ("Nákup: DM DROGERIE MARKT", -150, -650),
                 ("Nákup: BISTRO NA ROHU", -160, -520), ("Nákup: LEKARNA BENU", -90, -480),
                 ("BOLT.EU/O/2607", -120, -380), ("Nákup: ALZA.CZ", -300, -3500), ("Nákup: IKEA PRAHA", -400, -2400),
                 ("Nákup: CINEMA CITY", -300, -700), ("Nákup: TRAFIKA U MOSTU", -60, -250),
                 ("Nákup: ZARA", -600, -1800), ("Vklad z bankomatu? Ne - výběr ATM", -1000, -2000)]
        for _ in range(26):
            text, lo, hi = rnd.choice(shops)
            if text.startswith("Vklad"):
                text = "Výběr hotovosti ATM Praha"
            rows.append((d0 + timedelta(days=rnd.randint(0, 27)), -rnd.randint(-hi, -lo) if False else rnd.randint(hi, lo), text, ""))
    rows.sort(key=lambda x: x[0])
    for i, (d, amount, text, extra) in enumerate(rows):
        r = TX_ROWS[0] + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = d, amount, text, extra
    # Převod na spořicí účet: ruční kategorie, ukázka sloupce Vlastní kategorie
    for i, (_, _, text, _) in enumerate(rows):
        if text.startswith("Převod na spořicí"):
            ws[f"E{TX_ROWS[0] + i}"] = "Převod mezi účty"


def main():
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = "Přehled"
    sheets = {name: wb.create_sheet(name) for name in ["Výpis", "Kategorie", "Pravidla", "Návod"]}
    wb.defined_names["Rok"] = DefinedName("Rok", attr_text="'Přehled'!$C$4")
    wb.defined_names["Mesic"] = DefinedName("Mesic", attr_text="'Přehled'!$C$5")
    build_statement(sheets["Výpis"])
    build_categories(sheets["Kategorie"])
    rule_count = build_rules(sheets["Pravidla"])
    build_dashboard(ws_dash)
    build_guide(sheets["Návod"], rule_count)
    if DEMO:
        demo_transactions(sheets["Výpis"])
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = "Rodinný rozpočet"
    wb.properties.creator = "Kam to teče?"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT, f"({rule_count} pravidel)")


if __name__ == "__main__":
    main()
