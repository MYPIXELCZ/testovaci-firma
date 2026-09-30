#!/usr/bin/env python3
"""Generuje svatební plánovač Ano, beru (.xlsx).

    python3 product/build_planner.py          # prodejní verze (vzorové řádky označené „příklad“)
    python3 product/build_planner.py --demo   # vyplněná ukázka pro screenshoty na web
"""
import json
import sys
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DEMO = "--demo" in sys.argv
OUT = ROOT / "product" / "dist" / ("svatebni-planovac-demo.xlsx" if DEMO else "svatebni-planovac-anoberu.xlsx")

# Brand (brand/BRAND.md). Georgia/Arial: dostupné v Excelu i Google Tabulkách.
INK, MUTED, LINE = "2B2A28", "6B665E", "E6E0D6"
SAGE, SAGE_DARK, SAGE_LIGHT = "7C8B6F", "56654A", "EEF1EA"
TERRA, TERRA_DARK, CALC = "B5694A", "9A5238", "F3EFE8"
HEAD, BODY = "Georgia", "Arial"

F_TITLE = Font(name=HEAD, size=22, color=INK)
F_SECTION = Font(name=HEAD, size=13, bold=True, color=SAGE_DARK)
F_HINT = Font(name=BODY, size=9, italic=True, color=MUTED)
F_BODY = Font(name=BODY, size=10, color=INK)
F_LABEL = Font(name=BODY, size=10, color=MUTED)
F_BIG = Font(name=HEAD, size=16, bold=True, color=INK)
F_BAR = Font(name=BODY, size=9, color=SAGE)
F_HEAD = Font(name=BODY, size=10, bold=True, color="FFFFFF")
FILL_HEAD = PatternFill("solid", fgColor=SAGE_DARK)
FILL_INPUT = PatternFill("solid", fgColor=SAGE_LIGHT)
FILL_CALC = PatternFill("solid", fgColor=CALC)
B_ROW = Border(bottom=Side(style="thin", color=LINE))
B_INPUT = Border(*(Side(style="thin", color=SAGE),) * 4)

DATE_FMT = "d. m. yyyy"
KC_FMT = '#,##0 "Kč";-#,##0 "Kč";"–"'
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")

WEDDING = date(2027, 6, 12)  # demo datum (sobota)

# Orientační rozdělení rozpočtu podle běžné praxe (ne průzkum), uvedeno i v listu Návod.
CATEGORIES = [
    ("Místo a hostina", 0.40), ("Foto a video", 0.12), ("Oblečení a doplňky", 0.09),
    ("Prsteny", 0.07), ("Hudba a zábava", 0.07), ("Květiny a výzdoba", 0.07),
    ("Dort a sladkosti", 0.03), ("Vlasy a líčení", 0.02), ("Oznámení a tiskoviny", 0.02),
    ("Doprava", 0.02), ("Obřad a poplatky", 0.02), ("Dárky", 0.02), ("Rezerva", 0.05),
    ("Svatební cesta", 0.0),
]

# (fáze, úkol, dní před svatbou; záporné = po svatbě)
TASKS = [
    ("12+ měsíců", "Ujasnit si představu: velikost svatby, styl, roční období", 390),
    ("12+ měsíců", "Stanovit celkový rozpočet a domluvit, kdo co platí", 380),
    ("12+ měsíců", "Sepsat první seznam hostů (hrubý odhad počtu)", 370),
    ("12+ měsíců", "Vybrat termín a 2–3 náhradní", 365),
    ("12+ měsíců", "Vybrat a rezervovat místo obřadu i hostiny", 350),
    ("12+ měsíců", "Zjistit na matrice doklady, lhůty a poplatky pro váš termín a místo", 345),
    ("12+ měsíců", "Domluvit oddávajícího (matrika, církev, obřadník)", 340),
    ("12+ měsíců", "Vybrat a rezervovat fotografa", 330),
    ("12+ měsíců", "Vybrat a rezervovat kameramana", 320),
    ("9–12 měsíců", "Vybrat a požádat svědky", 300),
    ("9–12 měsíců", "Rezervovat kapelu nebo DJ", 290),
    ("9–12 měsíců", "Rezervovat catering (pokud není součástí místa)", 285),
    ("9–12 měsíců", "Poslat „save the date“ hostům, kteří cestují zdaleka", 280),
    ("9–12 měsíců", "Sjednotit styl svatby (barvy, výzdoba, nástěnka s inspirací)", 275),
    ("6–9 měsíců", "Začít vybírat svatební šaty", 250),
    ("6–9 měsíců", "Rezervovat floristku a výzdobu", 230),
    ("6–9 měsíců", "Rezervovat kadeřnici a vizážistku", 220),
    ("6–9 měsíců", "Zajistit ubytování pro hosty", 210),
    ("6–9 měsíců", "Naplánovat svatební cestu, zkontrolovat platnost pasů", 200),
    ("6–9 měsíců", "Vybrat oblek pro ženicha", 190),
    ("6–9 měsíců", "Rezervovat dopravu (svatební auto, autobus pro hosty)", 185),
    ("3–6 měsíců", "Objednat snubní prsteny", 160),
    ("3–6 měsíců", "Objednat svatební oznámení a pozvánky na hostinu", 150),
    ("3–6 měsíců", "Objednat svatební dort a koláčky", 130),
    ("3–6 měsíců", "Ochutnávka menu a výběr jídla", 110),
    ("3–6 měsíců", "Rozeslat oznámení a pozvánky", 100),
    ("3–6 měsíců", "Domluvit program: proslovy, hry, první tanec", 95),
    ("3–6 měsíců", "Zkouška účesu a líčení", 90),
    ("1–3 měsíce", "Podat na matrice dotazník k uzavření manželství s doklady", 75),
    ("1–3 měsíce", "Rozhodnout o příjmení po svatbě (uvádí se v dotazníku)", 75),
    ("1–3 měsíce", "Nacvičit první tanec", 60),
    ("1–3 měsíce", "Koupit doplňky: boty, závoj, kravata, podvazek", 50),
    ("1–3 měsíce", "Doplatit zálohy podle smluv s dodavateli", 45),
    ("1–3 měsíce", "Zkouška šatů a úpravy", 40),
    ("1–3 měsíce", "Připomenout se hostům, kteří neodpověděli", 35),
    ("1–3 měsíce", "Sepsat harmonogram dne D a poslat ho dodavatelům", 35),
    ("1–3 měsíce", "Nahlásit konečný počet hostů místu a cateringu", 30),
    ("1–3 měsíce", "Připravit zasedací pořádek", 30),
    ("2–4 týdny", "Vyzvednout snubní prsteny a zkontrolovat velikost", 21),
    ("2–4 týdny", "Připravit jmenovky, menu a tabuli se zasedacím pořádkem", 21),
    ("2–4 týdny", "Potvrdit časy se všemi dodavateli", 14),
    ("2–4 týdny", "Připravit obálky s doplatky pro dodavatele", 10),
    ("2–4 týdny", "Rozdělit úkoly na den D svědkům a pomocníkům", 10),
    ("Poslední týden", "Poslední zkouška šatů", 7),
    ("Poslední týden", "Sbalit nouzovou tašku (jehla a nit, náplasti, léky, deodorant)", 4),
    ("Poslední týden", "Odvézt výzdobu a věci na místo konání", 2),
    ("Poslední týden", "Odpočinout si a jít brzy spát", 1),
    ("Den D", "Občanské průkazy a prsteny má u sebe svědek", 0),
    ("Den D", "Najíst se a užít si to", 0),
    ("Po svatbě", "Vrátit půjčené věci (oblek, dekorace)", -7),
    ("Po svatbě", "Vyzvednout oddací list, pokud jste ho nedostali při obřadu", -7),
    ("Po svatbě", "Vyřídit nové doklady po změně příjmení (lhůtu ověřte na úřadě)", -7),
    ("Po svatbě", "Nahlásit změnu jména bance, pojišťovně a zaměstnavateli", -14),
    ("Po svatbě", "Poděkovat hostům a dodavatelům", -14),
    ("Po svatbě", "Vybrat fotky a objednat album", -60),
]

# (kategorie, položka, [demo: plán, skutečně, zaplaceno, splatnost dní před svatbou, dodavatel])
BUDGET = [
    ("Místo a hostina", "Pronájem prostor", 25000, 25000, 10000, 60, "Statek Na Kopci"),
    ("Místo a hostina", "Svatební oběd", 36000, None, 0, 14, "Statek Na Kopci"),
    ("Místo a hostina", "Večerní raut", 18000, None, 0, 14, "Statek Na Kopci"),
    ("Místo a hostina", "Nápoje", 20000, None, 0, 14, "Statek Na Kopci"),
    ("Foto a video", "Fotograf", 30000, 32000, 10000, 30, "Foto Jan Dvořák"),
    ("Foto a video", "Kameraman", 25000, None, 0, None, ""),
    ("Oblečení a doplňky", "Svatební šaty", 18000, 21500, 21500, None, "Salon Bílá"),
    ("Oblečení a doplňky", "Boty a doplňky nevěsty", 3000, None, 0, None, ""),
    ("Oblečení a doplňky", "Oblek ženicha", 8000, None, 0, None, ""),
    ("Oblečení a doplňky", "Boty a doplňky ženicha", 2500, None, 0, None, ""),
    ("Prsteny", "Snubní prsteny", 18000, None, 0, 45, ""),
    ("Hudba a zábava", "Kapela / DJ", 15000, 14000, 5000, 7, "DJ Martin"),
    ("Hudba a zábava", "Hudba na obřad", 3000, None, 0, None, ""),
    ("Květiny a výzdoba", "Kytice nevěsty", 2500, None, 0, None, ""),
    ("Květiny a výzdoba", "Korsáže a kytičky pro svědky", 1200, None, 0, None, ""),
    ("Květiny a výzdoba", "Výzdoba obřadu", 6000, None, 0, None, ""),
    ("Květiny a výzdoba", "Výzdoba stolů", 7000, None, 0, None, ""),
    ("Dort a sladkosti", "Svatební dort", 5000, None, 0, 7, ""),
    ("Dort a sladkosti", "Svatební koláčky", 3000, None, 0, None, ""),
    ("Vlasy a líčení", "Kadeřnice a vizážistka", 5000, None, 0, None, ""),
    ("Oznámení a tiskoviny", "Oznámení a pozvánky", 3500, 3200, 3200, None, "Tiskárna Papírek"),
    ("Oznámení a tiskoviny", "Jmenovky, menu, tabule", 1500, None, 0, None, ""),
    ("Doprava", "Svatební auto", 3000, None, 0, None, ""),
    ("Doprava", "Doprava hostů", 4000, None, 0, None, ""),
    ("Obřad a poplatky", "Poplatky matrice a oddávajícímu", 3000, None, 0, None, ""),
    ("Dárky", "Výslužka pro hosty", 3000, None, 0, None, ""),
    ("Dárky", "Dárky pro rodiče a svědky", 3000, None, 0, None, ""),
    ("Rezerva", "Nečekané výdaje", 12500, None, 0, None, ""),
    ("Svatební cesta", "Letenky a ubytování", None, None, 0, None, ""),
]

DAY_PLAN = [
    (time(8, 0), "Snídaně a příprava nevěsty (účes, líčení)", "Kadeřnice, vizážistka"),
    (time(10, 0), "Fotograf fotí přípravy", "Fotograf"),
    (time(11, 0), "Ženich přijíždí pro nevěstu", "Ženich, svědci"),
    (time(11, 30), "Odjezd na obřad", "Řidič"),
    (time(12, 0), "Svatební obřad", "Oddávající"),
    (time(12, 30), "Gratulace a skupinové focení", "Fotograf, svědci"),
    (time(13, 30), "Svatební oběd (rozbití talíře, krmení polévkou)", "Místo konání"),
    (time(15, 0), "Krájení dortu", ""),
    (time(15, 30), "Focení novomanželů", "Fotograf"),
    (time(17, 0), "Zábava a hry", "Svědci"),
    (time(19, 0), "První tanec", "Kapela / DJ"),
    (time(20, 0), "Večerní raut", "Catering"),
    (time(22, 0), "Házení kytice", ""),
    (time(0, 0), "Půlnoční překvapení, čepení nevěsty", ""),
    (time(2, 0), "Konec oslavy", ""),
]

DEMO_GUESTS = [
    # jméno, strana, skupina, osob, děti, pozváni na, odesláno, odpověď, strava, ubytování, stůl
    ("Marie a Josef Novákovi", "Nevěsta", "Rodina", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ne", 1),
    ("Eva a Tomáš Svobodovi", "Ženich", "Rodina", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ne", 1),
    ("Babička Anna Nováková", "Nevěsta", "Rodina", 1, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez lepku", "Ano", 2),
    ("Petr Novák s rodinou", "Nevěsta", "Rodina", 4, 2, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ano", 2),
    ("Jana Horáková s rodinou", "Nevěsta", "Rodina", 3, 1, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ne", 2),
    ("Děda Karel Svoboda", "Ženich", "Rodina", 1, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ano", 3),
    ("Lucie a Martin Svobodovi", "Ženich", "Rodina", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Vegetarián", "Ano", 3),
    ("Kateřina Dvořáková", "Ženich", "Rodina", 2, 0, "Obřad a hostina", "Ano", "Čeká", "", "Ne", 3),
    ("Tereza Malá", "Nevěsta", "Přátelé", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Vegan", "Ano", 4),
    ("Barbora Černá", "Nevěsta", "Přátelé", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ano", 4),
    ("Klára Veselá", "Nevěsta", "Přátelé", 1, 0, "Obřad a hostina", "Ano", "Čeká", "", "Ne", 4),
    ("Ondřej Král", "Ženich", "Přátelé", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ano", 5),
    ("Jakub Beneš", "Ženich", "Přátelé", 1, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez omezení", "Ano", 5),
    ("Michal Pokorný", "Ženich", "Přátelé", 2, 0, "Obřad a hostina", "Ano", "Odmítnuto", "", "Ne", None),
    ("Adam a Nela Růžičkovi", "Společní", "Přátelé", 2, 0, "Obřad a hostina", "Ano", "Potvrzeno", "Bez laktózy", "Ne", 5),
    ("Šárka Kučerová", "Nevěsta", "Kolegové", 1, 0, "Jen večer", "Ano", "Potvrzeno", "Bez omezení", "Ne", 6),
    ("Filip Marek", "Ženich", "Kolegové", 2, 0, "Jen večer", "Ano", "Čeká", "", "Ne", None),
    ("Tým z práce (Ženich)", "Ženich", "Kolegové", 4, 0, "Jen večer", "Ne", "", "", "Ne", None),
    ("Sousedé Procházkovi", "Nevěsta", "Ostatní", 2, 0, "Jen oznámení", "Ano", "", "", "Ne", None),
    ("Paní učitelka Hájková", "Nevěsta", "Ostatní", 1, 0, "Jen oznámení", "Ne", "", "", "Ne", None),
]

DEMO_VENDORS = [
    ("Místo a hostina", "Statek Na Kopci", "Lenka Horká", "+420 601 111 222", "info@example.cz", "", 99000, 10000, 60, "Ano", "Rezervováno"),
    ("Foto a video", "Foto Jan Dvořák", "Jan Dvořák", "+420 602 333 444", "jan@example.cz", "", 32000, 10000, 30, "Ano", "Rezervováno"),
    ("Hudba a zábava", "DJ Martin", "Martin Kos", "+420 603 555 666", "dj@example.cz", "", 14000, 5000, 7, "Ano", "Rezervováno"),
    ("Oblečení a doplňky", "Salon Bílá", "", "+420 604 777 888", "", "", 21500, 21500, None, "Ne", "Zaplaceno"),
    ("Foto a video", "Kamera Petr", "Petr Vrba", "+420 605 999 000", "", "", 28000, 0, None, "Ne", "Poptáno"),
]

DEMO_TABLES = [("Hlavní stůl", 10), ("Rodina nevěsty", 8), ("Rodina ženicha", 8), ("Přátelé nevěsty", 8),
               ("Přátelé ženicha", 8), ("Kolegové", 8), ("Děti", 6)]

TASK_ROWS = (5, 89)
BUD_ROWS = (5, 89)
GUEST_ROWS = (5, 254)
VEND_ROWS = (5, 54)
TABLE_ROWS = (5, 29)
DAY_ROWS = (5, 40)
CAT_ROWS = (5, 4 + len(CATEGORIES))


def rng(sheet, col, rows):
    return f"'{sheet}'!${col}${rows[0]}:${col}${rows[1]}"


def bar(expr):
    n = f"ROUND(MIN(1,MAX(0,{expr}))*20,0)"
    return f'=REPT("■",{n})&REPT("□",20-{n})'


def setup_sheet(ws, title, hint, widths, tab=SAGE):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws.row_dimensions[1].height = 34
    ws["A2"] = hint
    ws["A2"].font = F_HINT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def table(ws, headers, rows, calc_cols=(), formats=None, center_cols=()):
    """Hlavička v řádku 4 a naformátovaná prázdná oblast pro data."""
    formats = formats or {}
    for i, (col, name) in enumerate(headers):
        c = ws[f"{col}4"]
        c.value, c.font, c.fill, c.alignment = name, F_HEAD, FILL_HEAD, Alignment(wrap_text=True, vertical="center", horizontal="center" if col in center_cols else "left")
    ws.row_dimensions[4].height = 32
    for r in range(rows[0], rows[1] + 1):
        for col, _ in headers:
            c = ws[f"{col}{r}"]
            c.font, c.border = F_BODY, B_ROW
            c.alignment = CENTER if col in center_cols else Alignment(vertical="center")
            if col in calc_cols:
                c.fill = FILL_CALC
            if col in formats:
                c.number_format = formats[col]
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"


def validation(ws, formula, cols, rows, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    if strict:
        dv.errorTitle, dv.error = "Neplatná hodnota", "Vyberte hodnotu ze seznamu."
    for col in cols:
        dv.add(f"{col}{rows[0]}:{col}{rows[1]}")
    ws.add_data_validation(dv)


def cf(ws, area, formula, font=None, fill=None):
    ws.conditional_formatting.add(area, FormulaRule(formula=[formula], font=font, fill=fill, stopIfTrue=False))


def example_note(ws, row, col):
    ws[f"{col}{row}"] = "příklad – přepište nebo smažte"
    ws[f"{col}{row}"].font = F_HINT


# ---------------------------------------------------------------- Úkoly
def build_tasks(ws):
    setup_sheet(ws, "Úkoly", "Termíny se počítají z data svatby. Hotové označte ve sloupci Stav. "
                "Vlastní úkol přidáte do prázdného řádku: stačí napsat, kolik dní před svatbou má být hotový.",
                {"A": 16, "B": 60, "C": 11, "D": 13, "E": 14, "F": 14, "G": 34})
    table(ws, [("A", "Fáze"), ("B", "Úkol"), ("C", "Dní před svatbou"), ("D", "Termín"), ("E", "Stav"),
               ("F", "Kdo"), ("G", "Poznámka")], TASK_ROWS, calc_cols="D",
          formats={"D": DATE_FMT, "C": "0"}, center_cols="CDE")
    today = date.today()
    for i, (phase, task, days) in enumerate(TASKS):
        r = TASK_ROWS[0] + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = phase, task, days
        if DEMO:
            due = WEDDING - timedelta(days=days)
            ws[f"E{r}"] = "Hotovo" if due < today - timedelta(days=5) else ("Rozpracováno" if due < today + timedelta(days=20) else None)
            ws[f"F{r}"] = "Oba" if i % 3 == 0 else None
    for r in range(TASK_ROWS[0], TASK_ROWS[1] + 1):
        ws[f"D{r}"] = f'=IF(OR($C{r}="",NOT(ISNUMBER(DatumSvatby))),"",DatumSvatby-$C{r})'
        ws[f"H{r}"] = f'=IF(OR($E{r}="Hotovo",$D{r}=""),"",$D{r}+ROW()/100000)'  # klíč pro „nejbližší úkoly“
    ws.column_dimensions["H"].hidden = True
    validation(ws, '"Hotovo,Rozpracováno"', "E", TASK_ROWS)
    validation(ws, '"Nevěsta,Ženich,Oba,Svědci,Rodina"', "F", TASK_ROWS, strict=False)
    a = f"A{TASK_ROWS[0]}:G{TASK_ROWS[1]}"
    r0 = TASK_ROWS[0]
    cf(ws, a, f'$E{r0}="Hotovo"', font=Font(color=MUTED, strike=True))
    cf(ws, f"D{r0}:D{TASK_ROWS[1]}", f'AND(ISNUMBER($D{r0}),$D{r0}<TODAY(),$E{r0}<>"Hotovo")', font=Font(color=TERRA_DARK, bold=True))
    cf(ws, f"D{r0}:D{TASK_ROWS[1]}", f'AND(ISNUMBER($D{r0}),$D{r0}>=TODAY(),$D{r0}-TODAY()<=30,$E{r0}<>"Hotovo")', font=Font(bold=True))


# ---------------------------------------------------------------- Rozpočet
def build_budget(ws):
    setup_sheet(ws, "Rozpočet", "Napište plán. Jakmile znáte skutečnou cenu, doplňte ji a plánovač začne počítat s ní. "
                "Zálohy a platby zapisujte do sloupce Zaplaceno.",
                {"A": 20, "B": 30, "C": 12, "D": 12, "E": 12, "F": 13, "G": 13, "H": 12, "I": 20, "J": 28,
                 "K": 3, "L": 20, "M": 11, "N": 13, "O": 13, "P": 13, "Q": 13})
    table(ws, [("A", "Kategorie"), ("B", "Položka"), ("C", "Plán"), ("D", "Skutečná cena"), ("E", "Zaplaceno"),
               ("F", "Počítá se s"), ("G", "Zbývá doplatit"), ("H", "Splatnost"), ("I", "Dodavatel"),
               ("J", "Poznámka")], BUD_ROWS, calc_cols="FG",
          formats={c: KC_FMT for c in "CDEFG"} | {"H": DATE_FMT}, center_cols="H")
    for i, (cat, item, plan, actual, paid, due, vendor) in enumerate(BUDGET):
        r = BUD_ROWS[0] + i
        ws[f"A{r}"], ws[f"B{r}"] = cat, item
        if DEMO:
            ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"I{r}"] = plan, actual, paid or None, vendor or None
            if due is not None:
                ws[f"H{r}"] = WEDDING - timedelta(days=due)
        elif item == "Fotograf":
            ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"I{r}"] = 30000, 32000, 10000, "Foto Jan Dvořák"
            example_note(ws, r, "J")
    for r in range(BUD_ROWS[0], BUD_ROWS[1] + 1):
        ws[f"F{r}"] = f'=IF(AND($C{r}="",$D{r}=""),"",IF($D{r}<>"",$D{r},N($C{r})))'
        ws[f"G{r}"] = f'=IF($F{r}="","",MAX(0,$F{r}-N($E{r})))'
        ws[f"Z{r}"] = f'=IF(OR($H{r}="",N($G{r})<=0),"",$H{r}+ROW()/100000)'  # klíč pro „nejbližší platby“
    ws.column_dimensions["Z"].hidden = True
    cats = f"$L${CAT_ROWS[0]}:$L${CAT_ROWS[1]}"
    validation(ws, f"={cats}", "A", BUD_ROWS)
    r0, r1 = BUD_ROWS
    cf(ws, f"D{r0}:D{r1}", f"AND(N($D{r0})>0,N($C{r0})>0,$D{r0}>$C{r0})", font=Font(color=TERRA_DARK, bold=True))
    cf(ws, f"H{r0}:H{r1}", f'AND(ISNUMBER($H{r0}),$H{r0}<TODAY(),N($G{r0})>0)', font=Font(color=TERRA_DARK, bold=True))

    # Souhrn po kategoriích
    heads = [("L", "Kategorie"), ("M", "Doporučený podíl"), ("N", "Doporučeno"), ("O", "Plán"),
             ("P", "Počítá se s"), ("Q", "Rozdíl")]
    for col, name in heads:
        c = ws[f"{col}4"]
        c.value, c.font, c.fill, c.alignment = name, F_HEAD, FILL_HEAD, WRAP
    for i, (cat, share) in enumerate(CATEGORIES):
        r = CAT_ROWS[0] + i
        ws[f"L{r}"], ws[f"M{r}"] = cat, share
        ws[f"N{r}"] = f"=$M{r}*N(Rozpocet)"
        ws[f"O{r}"] = f"=SUMIFS($C${r0}:$C${r1},$A${r0}:$A${r1},$L{r})"
        ws[f"P{r}"] = f"=SUMIFS($F${r0}:$F${r1},$A${r0}:$A${r1},$L{r})"
        ws[f"Q{r}"] = f"=N{r}-P{r}"
        for col in "LMNOPQ":
            c = ws[f"{col}{r}"]
            c.font, c.border = F_BODY, B_ROW
            c.number_format = "0%" if col == "M" else KC_FMT
            if col in "NOPQ":
                c.fill = FILL_CALC
    t = CAT_ROWS[1] + 1
    ws[f"L{t}"] = "Celkem"
    ws[f"M{t}"] = f"=SUM(M{CAT_ROWS[0]}:M{CAT_ROWS[1]})"
    for col in "NOPQ":
        ws[f"{col}{t}"] = f"=SUM({col}{CAT_ROWS[0]}:{col}{CAT_ROWS[1]})"
    for col in "LMNOPQ":
        c = ws[f"{col}{t}"]
        c.font = Font(name=BODY, size=10, bold=True, color=INK)
        c.number_format = "0%" if col == "M" else KC_FMT
        c.border = Border(top=Side(style="thin", color=INK))
    ws[f"L{t + 2}"] = ("Doporučený podíl je orientační rozdělení podle běžné praxe českých svateb. "
                       "Klidně ho přepište podle sebe. Rozdíl = doporučeno − počítá se s.")
    ws[f"L{t + 2}"].font = F_HINT
    ws[f"L{t + 2}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"L{t + 2}:Q{t + 4}")
    cf(ws, f"Q{CAT_ROWS[0]}:Q{t}", f"Q{CAT_ROWS[0]}<0", font=Font(color=TERRA_DARK, bold=True))


# ---------------------------------------------------------------- Hosté
def build_guests(ws):
    setup_sheet(ws, "Hosté", "Jeden řádek = jedna pozvánka (domácnost). Do sloupce Osob napište, kolik lidí na ni přijde. "
                "Odpovědi zapisujte průběžně, přehled se přepočítá sám.",
                {"A": 30, "B": 11, "C": 11, "D": 7, "E": 8, "F": 17, "G": 11, "H": 12, "I": 13, "J": 11,
                 "K": 7, "L": 30, "M": 24})
    table(ws, [("A", "Jméno"), ("B", "Strana"), ("C", "Skupina"), ("D", "Osob"), ("E", "Z toho děti"),
               ("F", "Pozvaní na"), ("G", "Oznámení odesláno"), ("H", "Odpověď"), ("I", "Strava"),
               ("J", "Ubytování"), ("K", "Stůl"), ("L", "Adresa / kontakt"), ("M", "Poznámka")],
          GUEST_ROWS, center_cols="DEGHJK", formats={"D": "0", "E": "0", "K": "0"})
    rows = DEMO_GUESTS if DEMO else [("Jana a Petr Novákovi", "Nevěsta", "Rodina", 2, 0, "Obřad a hostina",
                                      "Ano", "Čeká", "Bez omezení", "Ne", None)]
    for i, g in enumerate(rows):
        r = GUEST_ROWS[0] + i
        for col, v in zip("ABCDEFGHIJK", g):
            ws[f"{col}{r}"] = v if v != "" else None
    if not DEMO:
        example_note(ws, GUEST_ROWS[0], "M")
    validation(ws, '"Nevěsta,Ženich,Společní"', "B", GUEST_ROWS)
    validation(ws, '"Rodina,Přátelé,Kolegové,Ostatní"', "C", GUEST_ROWS, strict=False)
    validation(ws, '"Obřad a hostina,Jen večer,Jen oznámení"', "F", GUEST_ROWS)
    validation(ws, '"Ano,Ne"', "GJ", GUEST_ROWS)
    validation(ws, '"Potvrzeno,Odmítnuto,Čeká"', "H", GUEST_ROWS)
    validation(ws, '"Bez omezení,Vegetarián,Vegan,Bez lepku,Bez laktózy,Jiné"', "I", GUEST_ROWS, strict=False)
    r0, r1 = GUEST_ROWS
    cf(ws, f"A{r0}:M{r1}", f'$H{r0}="Odmítnuto"', font=Font(color=MUTED, strike=True))
    cf(ws, f"H{r0}:H{r1}", f'$H{r0}="Potvrzeno"', font=Font(color=SAGE_DARK, bold=True))


# ---------------------------------------------------------------- Stoly
def build_tables(ws):
    setup_sheet(ws, "Stoly", "Číslo stolu přiřaďte hostům v listu Hosté. Tady uvidíte, kolik míst je obsazeno.",
                {"A": 8, "B": 28, "C": 11, "D": 11, "E": 11, "F": 14, "G": 3, "H": 30, "I": 12})
    table(ws, [("A", "Stůl"), ("B", "Popis"), ("C", "Kapacita"), ("D", "Obsazeno"), ("E", "Volno"), ("F", "Stav")],
          TABLE_ROWS, calc_cols="DEF", center_cols="ACDEF")
    g = GUEST_ROWS
    for i in range(10):
        ws[f"A{TABLE_ROWS[0] + i}"] = i + 1
    if DEMO:
        for i, (name, cap) in enumerate(DEMO_TABLES):
            ws[f"B{TABLE_ROWS[0] + i}"], ws[f"C{TABLE_ROWS[0] + i}"] = name, cap
    else:
        ws[f"B{TABLE_ROWS[0]}"], ws[f"C{TABLE_ROWS[0]}"] = "Hlavní stůl", 10
    for r in range(TABLE_ROWS[0], TABLE_ROWS[1] + 1):
        ws[f"D{r}"] = (f'=IF($A{r}="","",SUMIFS({rng("Hosté", "D", g)},{rng("Hosté", "K", g)},$A{r},'
                       f'{rng("Hosté", "H", g)},"<>Odmítnuto"))')
        ws[f"E{r}"] = f'=IF(OR($A{r}="",$C{r}=""),"",$C{r}-$D{r})'
        ws[f"F{r}"] = f'=IF(OR($A{r}="",$C{r}=""),"",IF($D{r}>$C{r},"Přeplněno",IF($D{r}=$C{r},"Plno","")))'
    r0 = TABLE_ROWS[0]
    cf(ws, f"E{r0}:F{TABLE_ROWS[1]}", f'$F{r0}="Přeplněno"', font=Font(color=TERRA_DARK, bold=True))
    summary = [
        ("Kapacita všech stolů", f"=SUM(C{r0}:C{TABLE_ROWS[1]})"),
        ("Potvrzení hosté (osob)", f'=SUMIFS({rng("Hosté", "D", g)},{rng("Hosté", "H", g)},"Potvrzeno")'),
        ("Potvrzení hosté bez stolu", f'=SUMIFS({rng("Hosté", "D", g)},{rng("Hosté", "H", g)},"Potvrzeno",'
                                     f'{rng("Hosté", "K", g)},"")'),
    ]
    for i, (label, formula) in enumerate(summary):
        r = 5 + i
        ws[f"H{r}"], ws[f"I{r}"] = label, formula
        ws[f"H{r}"].font, ws[f"I{r}"].font = F_LABEL, Font(name=BODY, size=11, bold=True, color=INK)
        ws[f"I{r}"].fill, ws[f"I{r}"].alignment = FILL_CALC, CENTER
    cf(ws, "I7", "I7>0", font=Font(color=TERRA_DARK, bold=True))


# ---------------------------------------------------------------- Dodavatelé
def build_vendors(ws):
    setup_sheet(ws, "Dodavatelé", "Kontakty, ceny a stav domluvy na jednom místě. Stav Zamítnuto použijte pro nabídky, které jste nevzali.",
                {"A": 20, "B": 24, "C": 18, "D": 17, "E": 24, "F": 20, "G": 12, "H": 12, "I": 12, "J": 10,
                 "K": 13, "L": 28})
    table(ws, [("A", "Kategorie"), ("B", "Firma / jméno"), ("C", "Kontaktní osoba"), ("D", "Telefon"),
               ("E", "E-mail"), ("F", "Web"), ("G", "Cena"), ("H", "Záloha"), ("I", "Splatnost zálohy"),
               ("J", "Smlouva"), ("K", "Stav"), ("L", "Poznámka")], VEND_ROWS,
          formats={"G": KC_FMT, "H": KC_FMT, "I": DATE_FMT}, center_cols="IJK")
    rows = DEMO_VENDORS if DEMO else [("Foto a video", "Foto Jan Dvořák", "Jan Dvořák", "+420 777 123 456",
                                       "jan@example.cz", "", 32000, 10000, None, "Ano", "Rezervováno")]
    for i, v in enumerate(rows):
        r = VEND_ROWS[0] + i
        for col, val in zip("ABCDEFGHIJK", v):
            if col == "I" and val is not None:
                val = WEDDING - timedelta(days=val)
            ws[f"{col}{r}"] = val if val != "" else None
    if not DEMO:
        example_note(ws, VEND_ROWS[0], "L")
    validation(ws, f"='Rozpočet'!$L${CAT_ROWS[0]}:$L${CAT_ROWS[1]}", "A", VEND_ROWS)
    validation(ws, '"Ano,Ne"', "J", VEND_ROWS)
    validation(ws, '"Poptáno,Rezervováno,Zaplaceno,Zamítnuto"', "K", VEND_ROWS)
    r0, r1 = VEND_ROWS
    cf(ws, f"A{r0}:L{r1}", f'$K{r0}="Zamítnuto"', font=Font(color=MUTED, strike=True))


# ---------------------------------------------------------------- Den D
def build_day(ws):
    setup_sheet(ws, "Den D", "Harmonogram svatebního dne. Upravte ho a pošlete fotografovi, kapele i svědkům.",
                {"A": 9, "B": 46, "C": 24, "D": 24, "E": 34})
    table(ws, [("A", "Čas"), ("B", "Co se děje"), ("C", "Kde"), ("D", "Kdo zajišťuje"), ("E", "Poznámka")],
          DAY_ROWS, formats={"A": "hh:mm"}, center_cols="A")
    for i, (t, what, who) in enumerate(DAY_PLAN):
        r = DAY_ROWS[0] + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"D{r}"] = t, what, who or None


# ---------------------------------------------------------------- Přehled
def build_dashboard(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TERRA
    for col, w in {"A": 2, "B": 30, "C": 16, "D": 8, "E": 4, "F": 32, "G": 14, "H": 12}.items():
        ws.column_dimensions[col].width = w
    ws["B2"] = "Svatební plánovač"
    ws["B2"].font = F_TITLE
    ws.row_dimensions[2].height = 36
    ws["B3"] = "Vyplňte zelená pole, zbytek se spočítá sám. Návod najdete v posledním listu."
    ws["B3"].font = F_HINT
    logo = Image(str(ROOT / "brand" / "logo@2x.png"))
    logo.width, logo.height = 184, 31
    ws.add_image(logo, "F2")

    def section(cell, text):
        ws[cell] = text
        ws[cell].font = F_SECTION
        ws[cell].border = Border(bottom=Side(style="thin", color=SAGE))
        col, row = cell[0], cell[1:]
        nxt = {"B": "CD", "F": "GH"}[col]
        for c in nxt:
            ws[f"{c}{row}"].border = Border(bottom=Side(style="thin", color=SAGE))

    def label(cell, text):
        ws[cell] = text
        ws[cell].font = F_LABEL
        ws[cell].alignment = Alignment(vertical="center")

    def value(cell, formula, fmt=None, big=False, merge_to=None):
        ws[cell] = formula
        ws[cell].font = F_BIG if big else Font(name=BODY, size=11, bold=True, color=INK)
        ws[cell].alignment = Alignment(horizontal="right", vertical="center")
        if fmt:
            ws[cell].number_format = fmt
        if merge_to:
            ws.merge_cells(f"{cell}:{merge_to}")

    # Základní údaje
    section("B5", "Základní údaje")
    inputs = [("B6", "Nevěsta", "Tereza" if DEMO else None, None),
              ("B7", "Ženich", "Jakub" if DEMO else None, None),
              ("B8", "Datum svatby", WEDDING if DEMO else None, DATE_FMT),
              ("B9", "Celkový rozpočet", 280000 if DEMO else None, KC_FMT)]
    for cell, text, v, fmt in inputs:
        r = cell[1:]
        label(cell, text)
        c = ws[f"C{r}"]
        c.value, c.fill, c.font = v, FILL_INPUT, Font(name=BODY, size=11, color=INK)
        c.alignment = Alignment(horizontal="left", vertical="center")
        if fmt:
            c.number_format = fmt
        for col in "CD":
            ws[f"{col}{r}"].border = B_INPUT
            ws[f"{col}{r}"].fill = FILL_INPUT
        ws.merge_cells(f"C{r}:D{r}")
        ws.row_dimensions[int(r)].height = 22

    # Odpočet
    section("F5", "Odpočet")
    label("F6", "Do svatby zbývá")
    value("G6", '=IF(ISNUMBER(DatumSvatby),MAX(0,DatumSvatby-TODAY()),"–")', "0", big=True)
    ws["H6"] = '=IF(ISNUMBER(G6),IF(G6=1,"den",IF(AND(G6>=2,G6<=4),"dny","dní")),"")'
    ws["H6"].font = F_LABEL
    label("F7", "Svatba připadá na")
    value("G7", '=IF(ISNUMBER(DatumSvatby),CHOOSE(WEEKDAY(DatumSvatby,2),"pondělí","úterý","středa",'
                '"čtvrtek","pátek","sobotu","neděli"),"")', merge_to="H7")
    t = TASK_ROWS
    label("F8", "Hotové úkoly")
    value("G8", f'=IFERROR(COUNTIF({rng("Úkoly", "E", t)},"Hotovo")/COUNTA({rng("Úkoly", "B", t)}),0)', "0%")
    ws["G9"] = bar("G8")
    ws["G9"].font = F_BAR
    ws.merge_cells("G9:H9")

    # Rozpočet
    b = BUD_ROWS
    section("B11", "Rozpočet")
    rows = [
        ("B12", "Počítá se s (plán nebo skutečná cena)", f'=SUM({rng("Rozpočet", "F", b)})'),
        ("B13", "Zbývá z rozpočtu", '=IF(N(Rozpocet)=0,"–",Rozpocet-C12)'),
        ("B14", "Už zaplaceno", f'=SUM({rng("Rozpočet", "E", b)})'),
        ("B15", "Zbývá doplatit", f'=SUM({rng("Rozpočet", "G", b)})'),
    ]
    for cell, text, formula in rows:
        label(cell, text)
        value(f"C{cell[1:]}", formula, KC_FMT, merge_to=f"D{cell[1:]}")
    label("B16", "Čerpání rozpočtu")
    value("C16", "=IF(N(Rozpocet)=0,0,C12/Rozpocet)", "0%", merge_to="D16")
    ws["C17"] = bar("C16")
    ws["C17"].font = F_BAR
    ws.merge_cells("C17:D17")
    cf(ws, "C13", "AND(ISNUMBER(C13),C13<0)", font=Font(color=TERRA_DARK, bold=True))
    cf(ws, "C16", "C16>1", font=Font(color=TERRA_DARK, bold=True))
    cf(ws, "C17", "C16>1", font=Font(color=TERRA))

    # Hosté
    g = GUEST_ROWS
    D, E, F, H, I, K = (rng("Hosté", c, g) for c in "DEFHIK")
    section("F11", "Hosté (počet osob)")
    rows = [
        ("F12", "Pozvaní na obřad a hostinu", f'=SUMIFS({D},{F},"Obřad a hostina")'),
        ("F13", "Pozvaní jen na večer", f'=SUMIFS({D},{F},"Jen večer")'),
        ("F14", "Potvrzeno", f'=SUMIFS({D},{H},"Potvrzeno")'),
        ("F15", "Odmítnuto", f'=SUMIFS({D},{H},"Odmítnuto")'),
        ("F16", "Čeká na odpověď", f'=SUMIFS({D},{F},"<>Jen oznámení",{F},"<>",{H},"<>Potvrzeno",{H},"<>Odmítnuto")'),
        ("F17", "Z potvrzených dětí", f'=SUMIFS({E},{H},"Potvrzeno")'),
        ("F18", "Z potvrzených se speciální stravou", f'=SUMIFS({D},{H},"Potvrzeno",{I},"<>Bez omezení",{I},"<>")'),
        ("F19", "Potvrzení bez stolu", f'=SUMIFS({D},{H},"Potvrzeno",{K},"")'),
    ]
    for cell, text, formula in rows:
        label(cell, text)
        value(f"G{cell[1:]}", formula, "0")
    cf(ws, "G19", "G19>0", font=Font(color=TERRA_DARK, bold=True))

    # Nejbližší úkoly a platby
    section("B21", "Nejbližší úkoly")
    section("F21", "Nejbližší platby")
    for cell, text in [("B22", "Úkol"), ("C22", "Termín"), ("F22", "Položka"), ("G22", "Splatnost"), ("H22", "Doplatit")]:
        ws[cell] = text
        ws[cell].font = Font(name=BODY, size=9, bold=True, color=MUTED)
    tk, tb = rng("Úkoly", "H", t), rng("Rozpočet", "Z", b)
    for k in range(1, 6):
        r = 22 + k
        done, no_date, no_pay = ('"Všechno hotovo"', '"Vyplňte datum svatby"', '"Žádné naplánované platby"') \
            if k == 1 else ('""', '""', '""')
        ws[f"B{r}"] = (f'=IFERROR(INDEX({rng("Úkoly", "B", t)},MATCH(SMALL({tk},{k}),{tk},0)),'
                       f'IF(ISNUMBER(DatumSvatby),{done},{no_date}))')
        ws[f"C{r}"] = f'=IFERROR(INT(SMALL({tk},{k})),"")'
        ws[f"F{r}"] = f'=IFERROR(INDEX({rng("Rozpočet", "B", b)},MATCH(SMALL({tb},{k}),{tb},0)),{no_pay})'
        ws[f"G{r}"] = f'=IFERROR(INT(SMALL({tb},{k})),"")'
        ws[f"H{r}"] = f'=IFERROR(INDEX({rng("Rozpočet", "G", b)},MATCH(SMALL({tb},{k}),{tb},0)),"")'
        for col in "BCFGH":
            c = ws[f"{col}{r}"]
            c.font, c.border = F_BODY, B_ROW
            c.alignment = Alignment(vertical="center", wrap_text=col in "BF")
        ws[f"C{r}"].number_format = ws[f"G{r}"].number_format = DATE_FMT
        ws[f"H{r}"].number_format = KC_FMT
        ws.merge_cells(f"C{r}:D{r}")
        ws.row_dimensions[r].height = 28
        cf(ws, f"C{r}", f"AND(ISNUMBER(C{r}),C{r}<TODAY())", font=Font(color=TERRA_DARK, bold=True))
        cf(ws, f"G{r}", f"AND(ISNUMBER(G{r}),G{r}<TODAY())", font=Font(color=TERRA_DARK, bold=True))
    ws["B29"] = "Červeně = po termínu."
    ws["B29"].font = F_HINT
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---------------------------------------------------------------- Návod
def build_guide(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = MUTED
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 90
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    lines = [
        ("title", "Jak plánovač používat"),
        ("text", ""),
        ("h", "Začněte tady"),
        ("text", "1. Na listu Přehled vyplňte jména, datum svatby a celkový rozpočet (zelená pole)."),
        ("text", "2. Úkoly: termíny se dopočítají podle data svatby. Hotové úkoly označte ve sloupci Stav."),
        ("text", "3. Rozpočet: ke každé položce napište plán. Jakmile znáte skutečnou cenu, doplňte ji a plánovač "
                 "začne počítat s ní."),
        ("text", "4. Hosté: jeden řádek = jedna pozvánka. Do sloupce Osob napište, kolik lidí na ni přijde."),
        ("text", "5. Stoly: v listu Hosté napište číslo stolu, v listu Stoly pak uvidíte obsazenost."),
        ("text", "6. Den D: upravte harmonogram a pošlete ho fotografovi, kapele i svědkům."),
        ("text", ""),
        ("h", "Tipy"),
        ("text", "• Béžově podbarvené sloupce se počítají samy, nepřepisujte je."),
        ("text", "• Řádky označené „příklad“ smažte nebo přepište."),
        ("text", "• Google Tabulky: soubor nahrajte na Disk Google a otevřete v Tabulkách. Pak ho můžete sdílet "
                 "s partnerem a upravovat ho oba zároveň."),
        ("text", "• Vlastní úkol, položku rozpočtu nebo hosta přidáte do kteréhokoli prázdného řádku."),
        ("text", ""),
        ("h", "Upozornění"),
        ("text", "Termíny úkolů a doporučené rozdělení rozpočtu jsou orientační, vycházejí z běžné praxe českých "
                 "svateb. Doklady, lhůty a poplatky pro sňatek vždy ověřte na své matrice."),
        ("text", ""),
        ("h", "Licence"),
        ("text", "Plánovač je určený pro osobní použití jednoho páru. Prosíme, nešiřte ho dál. "
                 "Děkujeme, že podporujete malou českou firmu."),
        ("text", ""),
        ("hint", "Ano, beru · anoberu.cz · verze 1.0"),
    ]
    for i, (kind, text) in enumerate(lines, start=2):
        c = ws[f"B{i}"]
        c.value = text or None
        c.font = {"title": F_TITLE, "h": F_SECTION, "hint": F_HINT}.get(kind, F_BODY)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            ws.row_dimensions[i].height = 34


def main():
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = "Přehled"
    sheets = {name: wb.create_sheet(name) for name in
              ["Úkoly", "Rozpočet", "Hosté", "Stoly", "Dodavatelé", "Den D", "Návod"]}
    wb.defined_names["DatumSvatby"] = DefinedName("DatumSvatby", attr_text="'Přehled'!$C$8")
    wb.defined_names["Rozpocet"] = DefinedName("Rozpocet", attr_text="'Přehled'!$C$9")
    build_dashboard(ws_dash)
    build_tasks(sheets["Úkoly"])
    build_budget(sheets["Rozpočet"])
    build_guests(sheets["Hosté"])
    build_tables(sheets["Stoly"])
    build_vendors(sheets["Dodavatelé"])
    build_day(sheets["Den D"])
    build_guide(sheets["Návod"])
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = "Svatební plánovač"
    wb.properties.creator = "Ano, beru"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)
    export_content()


def export_content():
    """Obsah plánovače pro web (články), aby web a produkt říkaly totéž."""
    out = ROOT / "web" / "src" / "content" / "planner.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    data = {
        "tasks": [{"phase": p, "task": t, "daysBefore": d} for p, t, d in TASKS],
        "categories": [{"name": n, "share": s} for n, s in CATEGORIES],
        "budgetItems": [{"category": c, "item": i} for c, i, *_ in BUDGET],
        "dayPlan": [{"time": t.strftime("%H:%M"), "what": w, "who": who} for t, w, who in DAY_PLAN],
    }
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
